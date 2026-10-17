#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""tools/build_docx.py — единый .docx нормативного комплекта со сквозной навигацией (Р-041).

Сборка: титул → содержание → Правила (со словарём терминов в 2.1) → Регламент AS-IS →
приложения А…Ф. Каждая wikilink-цель (заголовок, block-ID) получает закладку Word; ссылки
становятся внутренними гиперссылками (Ctrl+клик). Диаграммы ```mermaid отрисовываются в PNG
(mermaid-cli, кэш wrk/tmp/mermaid-cache) и вставляются картинками; при сбое рендера
блок вставляется кодом с предупреждением (Р-048). Верификация: каждая гиперссылка
имеет закладку.

Запуск:
  uv run --with python-docx python tools/build_docx.py --out dist/Комплект_правил.docx
  uv run --with python-docx python tools/build_docx.py --verify dist/Комплект_правил.docx
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
RU_APPENDIX_ORDER = "АБВГДЕЖИКЛМНПРСТУФ"

HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
FENCE = re.compile(r"^```")
FENCE_LANG = re.compile(r"^```([\w-]+)?\s*$")
TABLE = re.compile(r"^\s*\|")
SEP_ROW = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
LIST_RE = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")
QUOTE = re.compile(r"^>\s?(.*)$")
HR = re.compile(r"^\s*(?:---+|\*\*\*+)\s*$")
BLOCK_ID = re.compile(r"\s*\^([A-Za-z][\w-]*)\s*$")
ROW_BLOCK_ID = re.compile(r"\s*\^([\w-]+)\s*\|\s*$")
WIKILINK = re.compile(r"\[\[([^\[\]]+?)\]\]")
MD_LINK = re.compile(r"\[([^\[\]]+)\]\((https?://[^)\s]+)\)")
INLINE_CODE = re.compile(r"`([^`]+)`")
BOLD = re.compile(r"\*\*([^*]+?)\*\*")
ITALIC = re.compile(r"(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])")


@dataclass
class Segment:
    text: str = ""
    bold: bool = False
    italic: bool = False
    code: bool = False


@dataclass
class Link:
    anchor: str = ""  # имя закладки Word (внутренняя) либо "" при external
    label: list = field(default_factory=list)
    external: str | None = None


@dataclass
class Wiki:
    loc: str  # «Файл», «Файл#^ID», «Файл#Заголовок»
    label: str = ""


@dataclass
class Inline:
    parts: list = field(default_factory=list)


@dataclass
class Block:
    kind: str  # heading | para | table | list | quote | code | hr
    level: int = 0
    inline: Inline | None = None
    rows: list | None = None  # list[(is_header, cells, row_id)]
    block_id: str | None = None
    text: str = ""
    ordered: bool = False
    indent: int = 0
    lang: str | None = None  # язык fence-блока (```mermaid → "mermaid")


@dataclass
class DocKit:
    path: Path
    role: str
    prefix: str
    title: str = ""
    blocks: list = field(default_factory=list)
    heading_marks: dict = field(default_factory=dict)  # № -> закладка
    heading_names: dict = field(default_factory=dict)  # текст -> закладка
    blockid_marks: dict = field(default_factory=dict)  # ID -> закладка


# ---------------------------------------------------------------- разбор markdown


def split_cells(row: str) -> list[str]:
    core = row.strip()
    if core.startswith("|"):
        core = core[1:]
    if core.endswith("|"):
        core = core[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", core)]


def parse_inline(text: str) -> Inline:
    """Разбор строковой разметки: wikilink, внешняя ссылка, код, жирный, курсив."""
    out = Inline()

    def emit_plain(chunk: str):
        if not chunk:
            return
        marks = []
        for m in BOLD.finditer(chunk):
            marks.append((m.start(), m.end(), "b", m.group(1)))
        for m in ITALIC.finditer(chunk):
            marks.append((m.start(), m.end(), "i", m.group(1)))
        marks.sort()
        last, bold, ital = 0, False, False
        for s, e, kind, piece in marks:
            if s < last:
                continue
            if chunk[last:s]:
                out.parts.append(Segment(chunk[last:s], bold, ital))
            bold, ital = (True, ital) if kind == "b" else (bold, True)
            out.parts.append(Segment(piece, bold, ital))
            last = e
        if chunk[last:]:
            out.parts.append(Segment(chunk[last:], bold, ital))

    pos, occupied = 0, -1
    tokens = []
    for rx, kind in ((WIKILINK, "wiki"), (MD_LINK, "md"), (INLINE_CODE, "code")):
        for m in rx.finditer(text):
            tokens.append((m.start(), m.end(), kind, m))
    tokens.sort()
    for s, e, kind, m in tokens:
        if s < occupied:
            continue
        emit_plain(text[pos:s])
        if kind == "code":
            out.parts.append(Segment(m.group(1), code=True))
        elif kind == "md":
            out.parts.append(Link(label=[Segment(m.group(1))], external=m.group(2)))
        else:
            inner = m.group(1)
            parts = inner.split("|", 1)
            out.parts.append(
                Wiki(
                    loc=parts[0].strip(), label=parts[1] if len(parts) > 1 else parts[0]
                )
            )
        pos, occupied = e, e
    emit_plain(text[pos:])
    return out


def plain_text(inline: Inline | None) -> str:
    if inline is None:
        return ""
    parts = []
    for p in inline.parts:
        if isinstance(p, Segment):
            parts.append(p.text)
        elif isinstance(p, (Link, Wiki)):
            label = p.label if isinstance(p.label, list) else [Segment(str(p.label))]
            parts.append("".join(s.text for s in label if isinstance(s, Segment)))
    return "".join(parts)


def parse_md(text: str) -> list[Block]:
    blocks: list[Block] = []
    para_buf: list[str] = []
    table_buf: list[tuple[str, str | None]] = []
    code_buf: list[str] = []
    code_lang: str | None = None
    fenced = False

    def flush_para():
        if para_buf:
            joined = " ".join(para_buf)
            bid = None
            m = BLOCK_ID.search(joined)
            if m:
                bid = m.group(1)
                joined = joined[: m.start()].rstrip()
            blocks.append(Block("para", inline=parse_inline(joined), block_id=bid))
            para_buf.clear()

    def flush_table():
        if not table_buf:
            return
        tbl, saw_sep = [], False
        for raw, row_id in table_buf:
            if SEP_ROW.match(raw):
                saw_sep = True
                continue
            cells = [parse_inline(c) for c in split_cells(raw)]
            tbl.append((not saw_sep, cells, row_id))
        blocks.append(Block("table", rows=tbl))
        table_buf.clear()

    for line in text.splitlines():
        if FENCE.match(line):
            if fenced:
                blocks.append(Block("code", text="\n".join(code_buf), lang=code_lang))
                code_buf.clear()
                code_lang = None
            else:
                flush_para()
                flush_table()
                m = FENCE_LANG.match(line)
                code_lang = m.group(1) if m else None
            fenced = not fenced
            continue
        if fenced:
            code_buf.append(line)
            continue
        if TABLE.match(line):
            flush_para()
            m = ROW_BLOCK_ID.search(line)
            if m:
                table_buf.append((line[: m.start()].rstrip(), m.group(1)))
            else:
                table_buf.append((line, None))
            continue
        if table_buf:
            flush_table()
        m = HEADING.match(line)
        if m:
            flush_para()
            blocks.append(
                Block("heading", level=len(m.group(1)), inline=parse_inline(m.group(2)))
            )
            continue
        if HR.match(line):
            flush_para()
            blocks.append(Block("hr"))
            continue
        m = LIST_RE.match(line)
        if m:
            flush_para()
            blocks.append(
                Block(
                    "list",
                    inline=parse_inline(m.group(3)),
                    ordered=m.group(2).endswith("."),
                    indent=min(len(m.group(1)) // 2, 1),
                )
            )
            continue
        m = QUOTE.match(line)
        if m:
            flush_para()
            blocks.append(Block("quote", inline=parse_inline(m.group(1))))
            continue
        if not line.strip():
            flush_para()
            continue
        para_buf.append(line.strip())
    flush_para()
    flush_table()
    return blocks


def kit_docs(root: Path) -> list[DocKit]:
    std = root / "wrk" / "standard"
    plan = [
        (std / "Правила.md", "std"),
        (std / "Регламент AS-IS.md", "reg"),
    ]
    app_dir = std / "appendices"
    if app_dir.is_dir():
        files = {p.stem.split(" ", 1)[0]: p for p in app_dir.glob("*.md")}
        ordered = [l for l in RU_APPENDIX_ORDER if l in files] + sorted(
            k for k in files if k not in RU_APPENDIX_ORDER
        )
        plan += [(files[l], f"app{n}") for n, l in enumerate(ordered)]
    kits = []
    for path, prefix in plan:
        if not path.exists():
            continue
        kit = DocKit(path, prefix, prefix)
        kit.blocks = parse_md(path.read_text(encoding="utf-8"))
        for b in kit.blocks:
            if b.kind == "heading" and b.level == 1 and not kit.title:
                kit.title = plain_text(b.inline)
                break
        kits.append(kit)
    return kits


# ---------------------------------------------------------------- сборка docx


def ensure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (
        ("Heading 1", 16),
        ("Heading 2", 14),
        ("Heading 3", 13),
        ("Heading 4", 12),
    ):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
    try:
        doc.styles["Hyperlink"]
    except KeyError:
        st = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        st.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        st.font.underline = True


def add_page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE \\* MERGEFORMAT")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    paragraph._p.append(fld)


def bookmark_name_for(block_id: str) -> str:
    return "a_" + block_id.replace("-", "_")


def find_chrome() -> str | None:
    """Путь к Chrome для mermaid-cli: MMDC_CHROME/CHROME_PATH либо кэш puppeteer."""
    env = os.environ.get("MMDC_CHROME") or os.environ.get("CHROME_PATH")
    if env and Path(env).exists():
        return env
    base = Path.home() / ".cache" / "puppeteer" / "chrome"
    if base.is_dir():
        cands = sorted(base.glob("*/chrome-linux64/chrome")) or sorted(
            base.glob("*/chrome-linux/chrome")
        )
        if cands:
            return str(cands[-1])
    return None


def render_mermaid(code: str, cache_dir: Path) -> tuple[Path | None, str | None]:
    """PNG для mermaid-диаграммы; кэш по SHA-256 кода (Р-048).

    Возвращает (путь, None) при успехе и (None, диагностика) при сбое —
    вызывающий код тогда вставляет блок кодом с предупреждением.
    """
    digest = hashlib.sha256(code.encode("utf-8")).hexdigest()[:16]
    png = cache_dir / f"mmd-{digest}.png"
    if png.exists():
        return png, None
    try:
        cache_dir.mkdir(parents=True, exist_ok=True)
        mmd = cache_dir / f"mmd-{digest}.mmd"
        mmd.write_text(code, encoding="utf-8")
        cmd = os.environ.get("MMDC_BIN", "npx -y @mermaid-js/mermaid-cli").split()
        cmd += ["-i", str(mmd), "-o", str(png), "-b", "white", "-s", "2"]
        chrome = find_chrome()
        if chrome:
            cfg = cache_dir / "pptr.json"
            cfg.write_text(
                json.dumps({"executablePath": chrome, "args": ["--no-sandbox"]}),
                encoding="utf-8",
            )
            cmd += ["-p", str(cfg)]
        env = dict(os.environ, PUPPETEER_SKIP_DOWNLOAD="1")
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180, env=env)
        if png.exists() and png.stat().st_size > 0:
            return png, None
        return None, (proc.stderr or proc.stdout or "пустой вывод mmdc").strip()[-200:]
    except Exception as exc:  # noqa: BLE001 — диагностика в предупреждение сборки
        return None, f"{type(exc).__name__}: {exc}"


class Builder:
    def __init__(self, doc: Document, mermaid_cache: Path | None = None):
        self.doc = doc
        self.bmk = 0
        self.by_stem: dict[str, DocKit] = {}
        self.problems: list[str] = []
        self.heading_no: dict[int, int] = {}
        self.mermaid_cache = mermaid_cache
        self.warnings: list[str] = []

    def add_bookmark(self, paragraph, name: str):
        self.bmk += 1
        s = OxmlElement("w:bookmarkStart")
        s.set(qn("w:id"), str(self.bmk))
        s.set(qn("w:name"), name)
        e = OxmlElement("w:bookmarkEnd")
        e.set(qn("w:id"), str(self.bmk))
        paragraph._p.insert(0, s)
        paragraph._p.append(e)

    def make_run(self, seg: Segment, hyperlink_style: bool = False) -> OxmlElement:
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        if hyperlink_style:
            st = OxmlElement("w:rStyle")
            st.set(qn("w:val"), "Hyperlink")
            rPr.append(st)
        if seg.bold:
            rPr.append(OxmlElement("w:b"))
        if seg.italic:
            rPr.append(OxmlElement("w:i"))
        if seg.code:
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), "Courier New")
            fonts.set(qn("w:hAnsi"), "Courier New")
            rPr.append(fonts)
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), "18")  # 9 пт
            rPr.append(sz)
        if len(rPr):
            r.append(rPr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = seg.text
        r.append(t)
        return r

    def relate_external(self, url: str) -> str:
        return self.doc.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )

    def resolve_wiki(self, loc: str) -> str | None:
        stem, _, anchor = loc.partition("#")
        kit = self.by_stem.get(stem.strip())
        if not kit:
            self.problems.append(f"wikilink на неизвестный документ «{stem}»")
            return None
        if not anchor:
            return kit.heading_marks.get(0)
        if anchor.startswith("^"):
            name = kit.blockid_marks.get(anchor[1:])
            if not name:
                self.problems.append(f"якорь {anchor} отсутствует в «{stem}»")
            return name
        name = kit.heading_names.get(anchor.strip())
        if not name:
            self.problems.append(f"заголовок «{anchor}» отсутствует в «{stem}»")
        return name

    def hyperlink_el(
        self, anchor: str | None, external: str | None, label: str
    ) -> OxmlElement:
        hl = OxmlElement("w:hyperlink")
        if external:
            hl.set(qn("r:id"), self.relate_external(external))
        else:
            hl.set(qn("w:anchor"), anchor or "")
        for part in parse_inline(label).parts:
            if isinstance(part, Segment):
                hl.append(self.make_run(part, hyperlink_style=True))
        if not len(hl):
            hl.append(self.make_run(Segment(label), hyperlink_style=True))
        return hl

    def render_inline(self, paragraph, inline: Inline):
        for part in inline.parts:
            if isinstance(part, Segment):
                paragraph._p.append(self.make_run(part))
            elif isinstance(part, Link):
                hl = OxmlElement("w:hyperlink")
                if part.external:
                    hl.set(qn("r:id"), self.relate_external(part.external))
                else:
                    hl.set(qn("w:anchor"), part.anchor)
                for seg in part.label:
                    if isinstance(seg, Segment):
                        hl.append(self.make_run(seg, hyperlink_style=True))
                paragraph._p.append(hl)
            elif isinstance(part, Wiki):
                bm = self.resolve_wiki(part.loc)
                if not bm:
                    paragraph._p.append(self.make_run(Segment(part.label)))
                else:
                    paragraph._p.append(self.hyperlink_el(bm, None, part.label))

    def render_block(self, kit: DocKit, b: Block, first: bool = False):
        d = self.doc
        if b.kind == "heading":
            style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}.get(
                b.level, "Heading 4"
            )
            p = d.add_paragraph(style=style)
            if first:
                p.paragraph_format.page_break_before = True
            self.render_inline(p, b.inline)
            n = self.heading_no[id(kit)] = self.heading_no.get(id(kit), -1) + 1
            mark = kit.heading_marks.get(n)
            if mark:
                self.add_bookmark(p, mark)
            return
        if b.kind == "para":
            p = d.add_paragraph()
            self.render_inline(p, b.inline)
            if b.block_id:
                mark = kit.blockid_marks.get(b.block_id)
                if mark:
                    self.add_bookmark(p, mark)
            return
        if b.kind == "list":
            styles = (
                ("List Number", "List Number 2")
                if b.ordered
                else ("List Bullet", "List Bullet 2")
            )
            p = d.add_paragraph(style=styles[min(b.indent, 1)])
            self.render_inline(p, b.inline)
            return
        if b.kind == "quote":
            p = d.add_paragraph(style="Intense Quote")
            self.render_inline(p, b.inline)
            return
        if b.kind == "code":
            if b.lang == "mermaid" and self.mermaid_cache is not None:
                png, err = render_mermaid(b.text, self.mermaid_cache)
                if png is not None:
                    p = d.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(str(png), width=Cm(16))
                    cap = d.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cr = cap.add_run("Диаграмма (Mermaid)")
                    cr.italic = True
                    cr.font.size = Pt(9)
                    return
                self.warnings.append(
                    f"mermaid не отрисована: {err} — блок вставлен кодом"
                )
            for ln in b.text.splitlines() or [""]:
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p._p.append(self.make_run(Segment(ln if ln else " ", code=True)))
            return
        if b.kind == "hr":
            p = d.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "999999")
            pbdr.append(bottom)
            pPr.append(pbdr)
            return
        if b.kind == "table":
            ncols = max((len(cells) for _, cells, _ in b.rows or []), default=0)
            if not ncols:
                return
            tbl = d.add_table(rows=0, cols=ncols)
            tbl.style = "Table Grid"
            for is_header, cells, row_id in b.rows or []:
                row = tbl.add_row()
                for ci in range(ncols):
                    para = row.cells[ci].paragraphs[0]
                    if ci < len(cells):
                        self.render_inline(para, cells[ci])
                    if is_header:
                        for r in para.runs:
                            r.bold = True
                if row_id:
                    mark = kit.blockid_marks.get(row_id)
                    if mark:
                        self.add_bookmark(row.cells[0].paragraphs[0], mark)
            return


def build(
    root: Path, out: Path, title: str, version: str, report: list | None = None
) -> Path:
    kits = kit_docs(root)
    if not kits:
        raise SystemExit("build: комплект не найден (wrk/standard/)")
    doc = Document()
    ensure_styles(doc)
    b = Builder(doc, mermaid_cache=root / "wrk" / "tmp" / "mermaid-cache")

    # пасс 1: имена закладок по документам
    for kit in kits:
        h_no = 0
        for blk in kit.blocks:
            if blk.kind == "heading":
                name = f"h_{kit.prefix}_{h_no}"
                kit.heading_marks[h_no] = name
                kit.heading_names.setdefault(plain_text(blk.inline), name)
                h_no += 1
            elif blk.kind == "para" and blk.block_id:
                kit.blockid_marks[blk.block_id] = bookmark_name_for(blk.block_id)
            elif blk.kind == "table":
                for _, _, row_id in blk.rows or []:
                    if row_id:
                        kit.blockid_marks.setdefault(row_id, bookmark_name_for(row_id))
        b.by_stem[kit.path.stem] = kit

    # титульный лист
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(220)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(22)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().add_run("")
    r2 = p2.add_run(f"Редакция {version}. Собрано {date.today():%d.%m.%Y}")
    r2.font.size = Pt(13)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Навигация: содержание и внутренние гиперссылки (Ctrl+клик)")
    r3.font.size = Pt(11)
    r3.italic = True

    # содержание
    toc = doc.add_paragraph(style="Heading 1")
    toc.paragraph_format.page_break_before = True
    toc.add_run("Содержание")
    for kit in kits:
        entry = doc.add_paragraph()
        entry._p.append(
            b.hyperlink_el(
                kit.heading_marks.get(0, ""), None, kit.title or kit.path.stem
            )
        )
        for blk in kit.blocks:
            if blk.kind == "heading" and blk.level == 2:
                sub = doc.add_paragraph()
                sub.paragraph_format.left_indent = Cm(0.75)
                hname = kit.heading_names.get(plain_text(blk.inline))
                if hname:
                    sub._p.append(b.hyperlink_el(hname, None, plain_text(blk.inline)))

    # документы комплекта
    for kit in kits:
        first = True
        for blk in kit.blocks:
            b.render_block(kit, blk, first=first)
            first = False

    # колонтитул с номером страницы
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer_p)

    if b.problems:
        for prob in b.problems:
            print(f"build: ОШИБКА: {prob}", file=sys.stderr)
        raise SystemExit(1)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    n_links = len(doc.element.body.findall(f".//{W}hyperlink"))
    n_bm = len(doc.element.body.findall(f".//{W}bookmarkStart"))
    print(
        f"build: {out} — документов {len(kits)}, гиперссылок {n_links},"
        f" закладок {n_bm}, диаграмм-картинок {len(doc.inline_shapes)}"
    )
    for warn in b.warnings:
        print(f"build: ПРЕДУПРЕЖДЕНИЕ: {warn}", file=sys.stderr)
    if report is not None:
        report.extend(b.warnings)
    return out


def verify_docx(path: Path) -> list[str]:
    """Каждая внутренняя гиперссылка имеет закладку; отчёт проблем."""
    doc = Document(str(path))
    names = {el.get(W + "name") for el in doc.element.iter(W + "bookmarkStart")}
    anchors = [
        el.get(W + "anchor")
        for el in doc.element.iter(W + "hyperlink")
        if el.get(W + "anchor")
    ]
    problems = [f"гиперссылка без закладки: {a}" for a in anchors if a not in names]
    print(
        f"verify: {path.name} — внутренних гиперссылок {len(anchors)},"
        f" внешних {sum(1 for r in doc.part.rels.values() if r.is_external)},"
        f" закладок {len(names)}, проблем {len(problems)}"
    )
    for prob in problems:
        print(f"verify: ОШИБКА: {prob}")
    return problems


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    ap.add_argument("--out", type=Path, default=Path("dist/Комплект_правил.docx"))
    ap.add_argument(
        "--title",
        default="Корпоративные правила проектирования моделей данных (КМД и ЛМД) — нормативный комплект",
    )
    ap.add_argument("--version", default="0.8-д7.6 (рабочая редакция)")
    ap.add_argument("--verify", type=Path, metavar="DOCX")
    args = ap.parse_args(argv)
    if args.verify:
        return 1 if verify_docx(args.verify) else 0
    build(args.root, args.out, args.title, args.version)
    return 1 if verify_docx(args.out) else 0


if __name__ == "__main__":
    sys.exit(main())
