# -*- coding: utf-8 -*-
"""Тесты сборщика единого .docx tools/build_docx.py (Р-041: навигация закладками).

Запуск: uv run --with jsonschema --with python-docx python -m unittest discover -s tests
Тесты работают на временном комплекте (tmp); материалы проекта не изменяются.
"""

import importlib.util
import struct
import tempfile
import unittest
import zlib
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]

try:
    import docx  # noqa: F401

    HAS_PYTHON_DOCX = True
except ImportError:
    HAS_PYTHON_DOCX = False

if HAS_PYTHON_DOCX:
    from docx import Document

    from tests.test_nav import load_module as load_nav, make_kit

    def load_builder():
        import sys

        spec = importlib.util.spec_from_file_location(
            "build_docx_under_test", REPO_ROOT / "tools" / "build_docx.py"
        )
        mod = importlib.util.module_from_spec(spec)
        sys.modules["build_docx_under_test"] = mod  # нужно для dataclasses
        spec.loader.exec_module(mod)
        return mod


@unittest.skipUnless(
    HAS_PYTHON_DOCX, "python-docx не установлен (uv run --with python-docx)"
)
class BuildDocxTest(unittest.TestCase):
    """Сборка tmp-комплекта и верификация навигации в .docx."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.nav = load_nav()
        cls.root = Path(cls.tmp.name)
        make_kit(cls.root)
        cls.nav.cmd_anchors(cls.root)
        cls.nav.cmd_linkify(cls.root)
        # дополнения: код-блок, список, цитата, внешняя ссылка
        std = cls.root / "wrk" / "standard" / "Стандарт.md"
        std.write_text(
            std.read_text(encoding="utf-8")
            + "\n```\ncode LDM-040 fence\n```\n"
            + "\n- пункт один;\n- пункт два.\n\n1. нумерованный;\n2. второй.\n"
            + "\n> Цитата-примечание.\n\n[Внешний](https://example.com) и хвост.\n",
            encoding="utf-8",
        )
        cls.builder = load_builder()
        cls.out = Path(cls.tmp.name) / "kit.docx"
        cls.builder.build(
            cls.root, cls.out, title="Тестовый комплект", version="0.1-тест"
        )

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def doc_xml_items(self, tag):
        doc = Document(str(self.out))
        return list(doc.element.iter(tag))

    def test_file_created(self):
        self.assertTrue(self.out.exists() and self.out.stat().st_size > 5000)

    def test_verify_no_problems(self):
        self.assertEqual(self.builder.verify_docx(self.out), [])

    def test_bookmarks_for_ids(self):
        names = {
            el.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name")
            for el in self.doc_xml_items(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkStart"
            )
        }
        for bm in ("a_T_001", "a_LDM_040", "a_PR_001", "a_CDM_043", "a_GEN_005"):
            self.assertIn(bm, names)

    def test_internal_hyperlinks_resolve(self):
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        anchors = {
            el.get(W + "anchor")
            for el in self.doc_xml_items(W + "hyperlink")
            if el.get(W + "anchor")
        }
        names = {el.get(W + "name") for el in self.doc_xml_items(W + "bookmarkStart")}
        self.assertTrue(anchors, "внутренних гиперссылок нет")
        self.assertEqual(anchors - names, set())

    def test_external_hyperlink_present(self):
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
        doc = Document(str(self.out))
        rels = doc.part.rels
        ext = [
            r for r in rels.values() if r.is_external and "example.com" in r.target_ref
        ]
        self.assertTrue(ext)

    def test_toc_links_to_docs(self):
        text = "\n".join(p.text for p in Document(str(self.out)).paragraphs)
        self.assertIn("Содержание", text)
        self.assertIn("Стандарт", text)


def _png_1px() -> bytes:
    """Валидный PNG 1x1 (RGBA) — собирается программно, без фиксации байтов."""

    def chunk(ctype: bytes, data: bytes) -> bytes:
        body = ctype + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    scanline = b"\x00" + b"\xff\xff\xff\xff"
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(scanline))
        + chunk(b"IEND", b"")
    )


PNG_1PX = _png_1px()


@unittest.skipUnless(
    HAS_PYTHON_DOCX, "python-docx не установлен (uv run --with python-docx)"
)
class MermaidDocxTest(unittest.TestCase):
    """Диаграммы Mermaid в .docx — картинки (Р-048), а не код-блоки."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.nav = load_nav()
        cls.root = Path(cls.tmp.name)
        make_kit(cls.root)
        std = cls.root / "wrk" / "standard" / "Стандарт.md"
        std.write_text(
            std.read_text(encoding="utf-8")
            + "\n```mermaid\nerDiagram\n    A ||--o{ B : содержит\n```\n",
            encoding="utf-8",
        )
        cls.mod = load_builder()
        cls.png = Path(cls.tmp.name) / "fake.png"
        cls.png.write_bytes(PNG_1PX)
        cls.orig_render = cls.mod.render_mermaid
        cls.mod.render_mermaid = lambda code, cache: (cls.png, None)
        cls.out = Path(cls.tmp.name) / "mmd.docx"
        cls.mod.build(cls.root, cls.out, title="М", version="0-тест")

    @classmethod
    def tearDownClass(cls):
        cls.mod.render_mermaid = cls.orig_render
        cls.tmp.cleanup()

    def test_fence_language_captured(self):
        blocks = self.mod.parse_md("```mermaid\ngraph TD\n```\n")
        self.assertEqual(blocks[0].kind, "code")
        self.assertEqual(blocks[0].lang, "mermaid")
        plain = self.mod.parse_md("```\ncode\n```\n")
        self.assertIsNone(plain[0].lang)

    def test_picture_embedded(self):
        self.assertEqual(len(Document(str(self.out)).inline_shapes), 1)

    def test_mermaid_code_not_emitted(self):
        text = "\n".join(p.text for p in Document(str(self.out)).paragraphs)
        self.assertNotIn("erDiagram", text)


@unittest.skipUnless(
    HAS_PYTHON_DOCX, "python-docx не установлен (uv run --with python-docx)"
)
class MermaidFallbackTest(unittest.TestCase):
    """Сбой рендера Mermaid — блок вставляется кодом с предупреждением (Р-048)."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.nav = load_nav()
        cls.root = Path(cls.tmp.name)
        make_kit(cls.root)
        std = cls.root / "wrk" / "standard" / "Стандарт.md"
        std.write_text(
            std.read_text(encoding="utf-8")
            + "\n```mermaid\nerDiagram\n    A ||--o{ B : содержит\n```\n",
            encoding="utf-8",
        )
        cls.mod = load_builder()
        cls.orig_render = cls.mod.render_mermaid
        cls.mod.render_mermaid = lambda code, cache: (None, "mmdc недоступен")
        cls.out = Path(cls.tmp.name) / "mmd-fb.docx"
        cls.report = []
        cls.mod.build(cls.root, cls.out, title="М", version="0-тест", report=cls.report)

    @classmethod
    def tearDownClass(cls):
        cls.mod.render_mermaid = cls.orig_render
        cls.tmp.cleanup()

    def test_code_fallback(self):
        text = "\n".join(p.text for p in Document(str(self.out)).paragraphs)
        self.assertIn("erDiagram", text)

    def test_warning_reported(self):
        self.assertTrue(self.report and "mermaid" in self.report[0].lower())


@unittest.skipUnless(
    HAS_PYTHON_DOCX, "python-docx не установлен (uv run --with python-docx)"
)
class RenderMermaidCacheTest(unittest.TestCase):
    """Кэш рендера по SHA-256 кода диаграммы; сбой — (None, диагностика) (Р-048)."""

    def test_cached_second_call(self):
        mod = load_builder()
        with tempfile.TemporaryDirectory() as td:
            cache = Path(td)
            calls = []

            def fake_run(cmd, **kw):
                calls.append(cmd)
                out = Path(cmd[cmd.index("-o") + 1])
                out.write_bytes(PNG_1PX)

                class P:
                    returncode = 0
                    stderr = ""
                    stdout = ""

                return P()

            orig = mod.subprocess.run
            mod.subprocess.run = fake_run
            try:
                p1, e1 = mod.render_mermaid("graph TD; A-->B", cache)
                p2, e2 = mod.render_mermaid("graph TD; A-->B", cache)
            finally:
                mod.subprocess.run = orig
            self.assertIsNone(e1)
            self.assertIsNone(e2)
            self.assertEqual(p1, p2)
            self.assertTrue(p1 and p1.exists())
            self.assertEqual(len(calls), 1)

    def test_failure_diagnosis(self):
        mod = load_builder()
        with tempfile.TemporaryDirectory() as td:

            def boom(cmd, **kw):
                raise FileNotFoundError("npx нет")

            orig = mod.subprocess.run
            mod.subprocess.run = boom
            try:
                png, err = mod.render_mermaid("graph TD; A-->B", Path(td))
            finally:
                mod.subprocess.run = orig
            self.assertIsNone(png)
            self.assertIn("npx", err)


if __name__ == "__main__":
    unittest.main()
